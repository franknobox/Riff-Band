#!/usr/bin/env python3
"""Build the polished non-technical AI4MS product introduction.

Design preset: narrative_proposal.
Named overrides: Noto Sans SC for Chinese coverage; AI4MS navy/teal brand
colors for title and headings; custom cover and process-table components.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "08_nontechnical_intro" / "AI4MS_非技术产品介绍书.md"
OUTPUT = ROOT / "08_nontechnical_intro" / "AI4MS_非技术产品介绍书.docx"

# narrative_proposal token map + named brand overrides.
FONT = "Noto Sans SC"
PAGE_W = Inches(8.5)
PAGE_H = Inches(11)
MARGIN = Inches(1)
HEADER_FOOTER = Inches(0.492)
CONTENT_W_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}

NAVY = "16324F"
TEAL = "1C7C84"
TEAL_DARK = "16646B"
BLUE = "2E74B5"
INK = "23313D"
MUTED = "5B6870"
LIGHT_TEAL = "EAF3F5"
LIGHT_BLUE = "EEF4F8"
LIGHT_GRAY = "F4F6F9"
BORDER = "D7E1E7"
GOLD = "C99A2E"
WHITE = "FFFFFF"

PAGE_BREAK_HEADINGS = {
    "2. 用户如何完成一项研究",
}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, **kwargs: int) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge in ("top", "start", "bottom", "end"):
        if edge not in kwargs:
            continue
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(kwargs[edge]))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = BORDER, size: int = 4, val: str = "single") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), val)
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = TABLE_INDENT_DXA) -> None:
    total = sum(widths_dxa)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell, **CELL_MARGINS_DXA)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_run_font(run, *, size: float | None = None, bold: bool | None = None,
                 color: str | None = None, italic: bool | None = None,
                 font: str = FONT) -> None:
    run.font.name = font
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), FONT)
    r_pr.append(r_fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), TEAL_DARK)
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    new_run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


INLINE_RE = re.compile(r"(\*\*.*?\*\*|`.*?`|https?://[^\s]+)")


def add_inline(paragraph, text: str, *, color: str = INK, size: float | None = None) -> None:
    for piece in INLINE_RE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            run = paragraph.add_run(piece[2:-2])
            set_run_font(run, size=size, bold=True, color=color)
        elif piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            set_run_font(run, size=(size or 10.5), color=NAVY, font="DejaVu Sans Mono")
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), LIGHT_GRAY)
            run._element.get_or_add_rPr().append(shd)
        elif piece.startswith("http://") or piece.startswith("https://"):
            add_hyperlink(paragraph, piece, piece)
        else:
            run = paragraph.add_run(piece)
            set_run_font(run, size=size, color=color)


def set_style_font(style, font: str, size: float, color: str, bold: bool) -> None:
    style.font.name = font
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), font)


def setup_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    set_style_font(normal, FONT, 11, INK, False)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    h1 = doc.styles["Heading 1"]
    set_style_font(h1, FONT, 16, NAVY, True)
    h1.paragraph_format.space_before = Pt(18)
    h1.paragraph_format.space_after = Pt(10)
    h1.paragraph_format.keep_with_next = True
    h1.paragraph_format.keep_together = True

    h2 = doc.styles["Heading 2"]
    set_style_font(h2, FONT, 13, TEAL_DARK, True)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.keep_with_next = True
    h2.paragraph_format.keep_together = True

    h3 = doc.styles["Heading 3"]
    set_style_font(h3, FONT, 12, "1F4D78", True)
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.keep_with_next = True
    h3.paragraph_format.keep_together = True


def add_numbering(doc: Document, kind: str) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids or [0]) + 1
    num_id = max(num_ids or [0]) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    lvl.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    lvl.append(lvl_jc)
    p_pr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    # narrative_proposal: text at 0.375in, hanging about 0.194in.
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "280")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "290")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    lvl.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), FONT)
    r_pr.append(r_fonts)
    lvl.append(r_pr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    override.append(start_override)
    num.append(override)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_el)
    paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT


def clone_numbering_instance(doc: Document, base_num_id: int) -> int:
    numbering = doc.part.numbering_part.element
    base = next(x for x in numbering.findall(qn("w:num")) if int(x.get(qn("w:numId"))) == base_num_id)
    abstract_ref = base.find(qn("w:abstractNumId")).get(qn("w:val"))
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    new_id = max(num_ids or [0]) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), abstract_ref)
    num.append(ref)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    override.append(start_override)
    num.append(override)
    numbering.append(num)
    return new_id


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run = OxmlElement("w:r")
    field_rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), MUTED)
    field_rpr.append(color)
    field_run.append(field_rpr)
    field_run.extend([begin, instr, separate, text, end])
    paragraph._p.append(field_run)
    tail = paragraph.add_run(" 页")
    set_run_font(tail, size=9, color=MUTED)


def setup_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = PAGE_W
    section.page_height = PAGE_H
    section.top_margin = MARGIN
    section.right_margin = MARGIN
    section.bottom_margin = MARGIN
    section.left_margin = MARGIN
    section.header_distance = HEADER_FOOTER
    section.footer_distance = HEADER_FOOTER
    section.different_first_page_header_footer = True

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.text = ""
    table = header.add_table(rows=1, cols=2, width=Inches(6.5))
    set_table_geometry(table, [4680, 4680], indent_dxa=0)
    set_table_borders(table, val="nil")
    left = table.cell(0, 0).paragraphs[0]
    right = table.cell(0, 1).paragraphs[0]
    left.paragraph_format.space_after = Pt(2)
    right.paragraph_format.space_after = Pt(2)
    r = left.add_run("AI4MS")
    set_run_font(r, size=9, bold=True, color=TEAL_DARK)
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = right.add_run("非技术产品介绍书 · v0.3")
    set_run_font(r, size=8.5, color=MUTED)
    # Quiet header rule.
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    bottom = borders.find(qn("w:bottom"))
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "5")
    bottom.set(qn("w:color"), BORDER)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.text = ""
    ft = footer.add_table(rows=1, cols=2, width=Inches(6.5))
    set_table_geometry(ft, [4680, 4680], indent_dxa=0)
    set_table_borders(ft, val="nil")
    lp = ft.cell(0, 0).paragraphs[0]
    rp = ft.cell(0, 1).paragraphs[0]
    lp.paragraph_format.space_before = Pt(2)
    rp.paragraph_format.space_before = Pt(2)
    r = lp.add_run("2026-07-16 · 产品策划与开发指导")
    set_run_font(r, size=8.5, color=MUTED)
    add_page_number(rp)


def add_cover(doc: Document) -> None:
    bar = doc.add_table(rows=1, cols=1)
    set_table_geometry(bar, [CONTENT_W_DXA], indent_dxa=0)
    set_table_borders(bar, val="nil")
    set_cell_shading(bar.cell(0, 0), TEAL)
    p = bar.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(42)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("AI FOR MANAGEMENT SCIENCE")
    set_run_font(r, size=10, bold=True, color=TEAL_DARK)
    r.font.all_caps = True

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.05
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("AI4MS")
    set_run_font(r, size=34, bold=True, color=NAVY)
    r = p.add_run("\n从一个研究想法，\n到一项可信的管理科学研究")
    set_run_font(r, size=24, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(26)
    r = p.add_run("非技术产品介绍书")
    set_run_font(r, size=15, color=TEAL_DARK)

    add_callout(doc, "先弄清别人已经做了什么，再决定我们真正值得做什么。", accent=GOLD, fill="FBF7EB", size=13)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("适合读者")
    set_run_font(r, size=9.5, bold=True, color=MUTED)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("高校管理者 · 学院与实验室负责人 · 教师 · 研究生 · 科研合作方 · 产品与运营团队")
    set_run_font(r, size=10.5, color=INK)

    strip = doc.add_table(rows=1, cols=3)
    set_table_geometry(strip, [3120, 3120, 3120], indent_dxa=0)
    set_table_borders(strip, val="nil")
    labels = [("找得全", NAVY), ("说得清", TEAL), ("有据可查", TEAL_DARK)]
    for idx, (label, fill) in enumerate(labels):
        cell = strip.cell(0, idx)
        set_cell_shading(cell, fill)
        set_cell_margins(cell, top=130, bottom=130, start=100, end=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(label)
        set_run_font(r, size=10.5, bold=True, color=WHITE)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(22)
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("版本 0.3 · 2026 年 7 月 16 日")
    set_run_font(r, size=9.5, color=MUTED)
    doc.add_page_break()


def add_callout(doc: Document, text: str, *, accent: str = TEAL, fill: str = LIGHT_TEAL,
                size: float = 11.5) -> None:
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [180, 9180])
    set_table_borders(table, val="nil")
    set_cell_shading(table.cell(0, 0), accent)
    set_cell_shading(table.cell(0, 1), fill)
    set_cell_margins(table.cell(0, 0), top=80, bottom=80, start=0, end=0)
    set_cell_margins(table.cell(0, 1), top=150, bottom=150, start=180, end=180)
    p = table.cell(0, 1).paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.22
    add_inline(p, text, color=NAVY, size=size)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_journey_table(doc: Document) -> None:
    cap = doc.add_paragraph()
    cap.paragraph_format.space_before = Pt(4)
    cap.paragraph_format.space_after = Pt(6)
    r = cap.add_run("研究旅程概览")
    set_run_font(r, size=10, bold=True, color=MUTED)

    rows = [
        ("01 说出想法", "与选题智能体把模糊方向整理成课题简报"),
        ("02 查看研究", "与文献智能体多源搜索、筛选并寻找反向证据"),
        ("03 人工选题", "比较候选问题，由研究者或导师完成 G0 审批"),
        ("04 理论与设计", "明确机制、识别或求解思路并完成 G1 审批"),
        ("05 数据与合规", "确认来源、许可、隐私和变量，完成 G2 审批"),
        ("06 计划与代码", "共同编辑分析计划或 Stata do-file，完成 G3 审批"),
        ("07 运行与复现", "保存环境、代码、日志和结果，执行稳健性检查"),
        ("08 证据与解释", "把主张连接到结果，由人工完成 G4 审批"),
        ("09 写作与发布", "从批准内容形成研究包，经 G5 审批后交付"),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    set_table_geometry(table, [2260, 7100])
    set_table_borders(table, color=BORDER, size=4, val="single")
    for idx, (label, desc) in enumerate(rows):
        left, right = table.rows[idx].cells
        set_cell_shading(left, NAVY if idx % 2 == 0 else TEAL_DARK)
        set_cell_shading(right, WHITE if idx % 2 == 0 else LIGHT_BLUE)
        lp = left.paragraphs[0]
        rp = right.paragraphs[0]
        lp.paragraph_format.space_after = Pt(0)
        rp.paragraph_format.space_after = Pt(0)
        r = lp.add_run(label)
        set_run_font(r, size=9.5, bold=True, color=WHITE)
        add_inline(rp, desc, color=INK, size=10)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.page_break_before = text in PAGE_BREAK_HEADINGS
    r = p.add_run(text)
    if level == 1:
        set_run_font(r, size=16, bold=True, color=NAVY)
    elif level == 2:
        set_run_font(r, size=13, bold=True, color=TEAL_DARK)
    else:
        set_run_font(r, size=12, bold=True, color="1F4D78")


def parse_markdown(doc: Document, markdown: str, bullet_num_id: int, decimal_num_id: int) -> None:
    lines = markdown.splitlines()
    start = next(i for i, line in enumerate(lines) if line.strip() == "## 一页读懂 AI4MS")
    lines = lines[start:]
    journey_inserted = False
    reference_section = False
    current_decimal_id = decimal_num_id
    previous_was_numbered = False
    i = 0
    while i < len(lines):
        raw = lines[i].rstrip()
        stripped = raw.strip()
        if not stripped or stripped == "---":
            j = i + 1
            while j < len(lines) and not lines[j].strip():
                j += 1
            next_is_numbered = j < len(lines) and re.match(r"^\d+\.\s+", lines[j].strip())
            if not (previous_was_numbered and next_is_numbered):
                previous_was_numbered = False
            i += 1
            continue
        if stripped.startswith(">"):
            block = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                block.append(lines[i].strip()[1:].strip())
                i += 1
            add_callout(doc, " ".join(block), accent=TEAL, fill=LIGHT_TEAL, size=12)
            previous_was_numbered = False
            continue
        if stripped.startswith("## "):
            title = stripped[3:].strip()
            add_heading(doc, title, 1)
            reference_section = title == "公开参考资料"
            previous_was_numbered = False
            i += 1
            continue
        if stripped.startswith("### "):
            title = stripped[4:].strip()
            if title == "第一步：说出想法" and not journey_inserted:
                add_journey_table(doc)
                journey_inserted = True
            add_heading(doc, title, 2)
            previous_was_numbered = False
            i += 1
            continue
        if stripped.startswith("#### "):
            add_heading(doc, stripped[5:].strip(), 3)
            previous_was_numbered = False
            i += 1
            continue
        bullet = re.match(r"^-\s+(.*)$", stripped)
        if bullet:
            p = doc.add_paragraph()
            apply_numbering(p, bullet_num_id)
            p.paragraph_format.space_after = Pt(2 if reference_section else 4)
            p.paragraph_format.line_spacing = 1.12 if reference_section else 1.208
            add_inline(p, bullet.group(1), color=INK, size=9.2 if reference_section else 10.8)
            previous_was_numbered = False
            i += 1
            continue
        numbered = re.match(r"^\d+\.\s+(.*)$", stripped)
        if numbered:
            if not previous_was_numbered:
                current_decimal_id = clone_numbering_instance(doc, decimal_num_id)
            p = doc.add_paragraph()
            apply_numbering(p, current_decimal_id)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.208
            add_inline(p, numbered.group(1), color=INK, size=10.8)
            previous_was_numbered = True
            i += 1
            continue

        # Collect wrapped prose until the next structural line.
        paragraph_lines = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt:
                break
            if nxt == "---" or nxt.startswith("#") or nxt.startswith(">"):
                break
            if re.match(r"^-\s+", nxt) or re.match(r"^\d+\.\s+", nxt):
                break
            paragraph_lines.append(nxt)
            i += 1
        text = " ".join(paragraph_lines)
        previous_was_numbered = False
        p = doc.add_paragraph()
        if text.startswith("说明："):
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = 1.2
            add_inline(p, text, color=MUTED, size=9.3)
        else:
            add_inline(p, text, color=INK, size=11)


def add_document_end_note(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("AI4MS · 让每一个研究决定有依据，让每一个研究结论可追溯")
    set_run_font(r, size=9.5, bold=True, color=TEAL_DARK)


def build() -> None:
    doc = Document()
    doc.core_properties.title = "AI4MS：从一个研究想法，到一项可信的管理科学研究"
    doc.core_properties.subject = "AI4MS 非技术产品介绍书"
    doc.core_properties.author = "AI4MS"
    doc.core_properties.keywords = "AI for Science, 管理科学, 科研平台, 课题侦察, 相关研究报告"
    doc.core_properties.comments = "Version 0.3, 2026-07-16"

    setup_styles(doc)
    setup_page(doc)
    bullet_num_id = add_numbering(doc, "bullet")
    decimal_num_id = add_numbering(doc, "decimal")
    add_cover(doc)
    parse_markdown(doc, SOURCE.read_text(encoding="utf-8"), bullet_num_id, decimal_num_id)
    add_document_end_note(doc)

    # Keep the source document in a single Letter section.
    for section in doc.sections:
        section.start_type = WD_SECTION.NEW_PAGE

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
